import docx
import pandas as pd
import os
import re

def clean_val(val):
    if not val: return 0
    # Remove dots used as thousands separators if any, and convert to int
    clean = re.sub(r'[^\d-]', '', str(val))
    try:
        return int(clean)
    except:
        return 0

def get_mes(etapa):
    mapping = {
        '1': 'Enero',
        '2': 'Febrero',
        '3': 'Marzo',
        '4': 'Abril',
        '5': 'Mayo',
        '6': 'Junio',
        '7': 'Julio',
        '8': 'Agosto',
        '9': 'Septiembre',
        '10': 'Octubre',
        '11': 'Noviembre',
        '12': 'Diciembre'
    }
    return mapping.get(str(etapa), etapa)

def process_carreras(file_path):
    doc = docx.Document(file_path)
    table = doc.tables[0]
    data = []
    for i, row in enumerate(table.rows):
        if i < 4: continue # Skip headers
        cells = [cell.text.strip() for cell in row.cells]
        if not cells[0] or cells[2].lower().startswith('total'): continue
        
        row_data = {
            'Etapa': cells[0],
            'Mes': get_mes(cells[0]),
            'Programa': cells[2],
            'Egresados': clean_val(cells[3]),
            'No matriculados': clean_val(cells[4]),
            'Admitidos': clean_val(cells[5]),
            'Admitidos Matriculados': clean_val(cells[6]),
            'Recuperados': clean_val(cells[7]),
            'Matríc. regular a Distancia': clean_val(cells[8]),
            'Matríc. regular Presencial': clean_val(cells[9]),
            'Estudiantes matrí. TOTAL': clean_val(cells[10]),
            'Alumnos en Asignaturas de 2 Meses': clean_val(cells[11])
        }
        data.append(row_data)
    
    df = pd.DataFrame(data)
    
    # Calculations
    df['Pérdida de Conversión Comercial(Fuga: Admit. - Adm. Mat.)'] = df['Admitidos'] - df['Admitidos Matriculados']
    df['Deserción Neta(Irreversible: No mat. - Recup.)'] = df['No matriculados'] - df['Recuperados']
    df['Brecha de Reemplazo(Vacío: Egresados - Adm. Mat.)'] = df['Egresados'] - df['Admitidos Matriculados']
    df['Crecimiento Neto Estudiantil(Balance: Entradas - Salidas)'] = (df['Admitidos Matriculados'] + df['Recuperados']) - (df['No matriculados'] + df['Egresados'])
    
    df['Tasa Efectividad Comercial(% Cierres: Adm.Mat./Admit.)'] = (df['Admitidos Matriculados'] / df['Admitidos']).fillna(0) * 100
    df['Tasa Éxito Retención(% Hemorragia: Rec./No mat.)'] = (df['Recuperados'] / df['No matriculados']).replace([float('inf'), -float('inf')], 0).fillna(0) * 100
    df['Tasa Renovación Genuina(% Inyección: Adm.Mat./Dist.)'] = (df['Admitidos Matriculados'] / df['Matríc. regular a Distancia']).replace([float('inf'), -float('inf')], 0).fillna(0) * 100
    df['Proporción de Matrícula en Cierre(% Presencial: Pres./Total)'] = (df['Matríc. regular Presencial'] / df['Estudiantes matrí. TOTAL']).fillna(0) * 100
    
    return df

def process_maestrias(file_path):
    doc = docx.Document(file_path)
    table = doc.tables[0]
    data = []
    for i, row in enumerate(table.rows):
        if i < 4: continue # Skip headers
        cells = [cell.text.strip() for cell in row.cells]
        # Maestrias has 13 columns, index 3 is Program
        if not cells[0] or cells[3].lower().startswith('total'): continue
        
        row_data = {
            'Etapa': cells[0],
            'Mes': get_mes(cells[0]),
            'Programa': cells[3],
            'Egresados': clean_val(cells[4]),
            'No matriculados': clean_val(cells[5]),
            'Admitidos': clean_val(cells[6]),
            'Admitidos Matriculados': clean_val(cells[7]),
            'Recuperados': clean_val(cells[8]),
            'Matríc. regular a Distancia': clean_val(cells[9]),
            'Matríc. regular Presencial': 0, # Posgrado has no Presencial
            'Estudiantes matrí. TOTAL': clean_val(cells[10]),
            'Alumnos en Asignaturas de 2 Meses': clean_val(cells[11])
        }
        data.append(row_data)
    
    df = pd.DataFrame(data)
    
    # Calculations
    df['Pérdida de Conversión Comercial(Fuga: Admit. - Adm. Mat.)'] = df['Admitidos'] - df['Admitidos Matriculados']
    df['Deserción Neta(Irreversible: No mat. - Recup.)'] = df['No matriculados'] - df['Recuperados']
    df['Brecha de Reemplazo(Vacío: Egresados - Adm. Mat.)'] = df['Egresados'] - df['Admitidos Matriculados']
    df['Crecimiento Neto Estudiantil(Balance: Entradas - Salidas)'] = (df['Admitidos Matriculados'] + df['Recuperados']) - (df['No matriculados'] + df['Egresados'])
    
    df['Tasa Efectividad Comercial(% Cierres: Adm.Mat./Admit.)'] = (df['Admitidos Matriculados'] / df['Admitidos']).fillna(0) * 100
    df['Tasa Éxito Retención(% Hemorragia: Rec./No mat.)'] = (df['Recuperados'] / df['No matriculados']).replace([float('inf'), -float('inf')], 0).fillna(0) * 100
    df['Tasa Renovación Genuina(% Inyección: Adm.Mat./Dist.)'] = (df['Admitidos Matriculados'] / df['Matríc. regular a Distancia']).replace([float('inf'), -float('inf')], 0).fillna(0) * 100
    df['Proporción de Matrícula en Cierre(% Presencial: Pres./Total)'] = 0 # Not applicable for Posgrado
    
    return df

# File paths
carreras_path = r"c:\Users\Renato Bertolotti\Documents\an-lisis-de-crecimiento-ia\Resultados Generales Carreras Profesionales MAY26ob.docx"
maestrias_path = r"c:\Users\Renato Bertolotti\Documents\an-lisis-de-crecimiento-ia\Resultados Generales Maestrías MAY26ob.docx"

output_carreras = r"c:\Users\Renato Bertolotti\Documents\an-lisis-de-crecimiento-ia\Cuadro_Mando_Pregrado_Actualizado.csv"
output_maestrias = r"c:\Users\Renato Bertolotti\Documents\an-lisis-de-crecimiento-ia\Cuadro_Mando_Posgrado_Actualizado.csv"

# Process and Save
print("Processing Carreras...")
df_carreras = process_carreras(carreras_path)
df_carreras.to_csv(output_carreras, index=False, encoding='utf-8-sig')
print(f"Saved to {output_carreras}")

print("Processing Maestrías...")
df_maestrias = process_maestrias(maestrias_path)
df_maestrias.to_csv(output_maestrias, index=False, encoding='utf-8-sig')
print(f"Saved to {output_maestrias}")
